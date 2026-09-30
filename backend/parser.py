from pathlib import Path
import re

import fitz
import pytesseract
from PIL import Image
from docx import Document


# Tesseract installation path on Windows
import shutil

tesseract_path = shutil.which("tesseract")

if tesseract_path:
    pytesseract.pytesseract.tesseract_cmd = tesseract_path
else:
    windows_tesseract = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

    if Path(windows_tesseract).exists():
        pytesseract.pytesseract.tesseract_cmd = windows_tesseract


SKILL_KEYWORDS = [
    "Python", "Java", "C", "C++", "C#", "JavaScript", "TypeScript",
    "HTML", "CSS", "React", "Angular", "Node.js",
    "FastAPI", "Django", "Flask",
    "SQL", "MySQL", "PostgreSQL", "MongoDB", "SQLite",
    "Git", "GitHub", "Docker",
    "AWS", "Azure", "GCP",
    "Power BI", "Tableau", "Excel",
    "Pandas", "NumPy", "scikit-learn",
    "TensorFlow", "PyTorch",
    "Machine Learning", "Deep Learning", "NLP",
    "REST API"
]


# --------------------------------------------------
# PDF TEXT EXTRACTION
# --------------------------------------------------

def extract_pdf_text(path):
    with fitz.open(path) as pdf:
        return "\n".join(page.get_text() for page in pdf)


# --------------------------------------------------
# PDF OCR
# --------------------------------------------------

def extract_pdf_ocr(path):
    text_parts = []

    with fitz.open(path) as pdf:

        for page in pdf:

            # Render PDF page as image
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))

            image = Image.frombytes(
                "RGB",
                [pix.width, pix.height],
                pix.samples
            )

            # OCR
            text = pytesseract.image_to_string(image)

            if text.strip():
                text_parts.append(text)

    return "\n".join(text_parts)


# --------------------------------------------------
# DOCX EXTRACTION
# --------------------------------------------------

def extract_docx(path):

    doc = Document(path)

    lines = [
        p.text
        for p in doc.paragraphs
        if p.text.strip()
    ]

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text)

    return "\n".join(lines)


# --------------------------------------------------
# GENERAL TEXT EXTRACTION
# --------------------------------------------------

def extract_text(path):

    if path.suffix.lower() == ".pdf":

        # First try normal PDF text extraction
        text = extract_pdf_text(path)

        # If no text found, use OCR
        if len(text.strip()) < 50:
            print("⚠️ No usable PDF text found. Running OCR...")
            text = extract_pdf_ocr(path)

        return text

    if path.suffix.lower() == ".docx":
        return extract_docx(path)

    raise ValueError("Unsupported file format.")


# --------------------------------------------------
# EMAIL
# --------------------------------------------------

def extract_email(text):

    pattern = r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b"

    match = re.search(pattern, text, re.I)

    return match.group(0) if match else ""


# --------------------------------------------------
# PHONE
# --------------------------------------------------

def extract_phone(text):

    patterns = [
        r"(?<!\d)(?:\+?\d{1,3}[-.\s]?)?\d{3}[-.\s]?\d{3}[-.\s]?\d{4}(?!\d)",
        r"(?<!\d)(?:\+?\d{1,3}[-.\s]?)?\d{10}(?!\d)"
    ]

    for pattern in patterns:

        match = re.search(pattern, text)

        if match:
            return match.group(0).strip()

    return ""


# --------------------------------------------------
# NAME
# --------------------------------------------------

def extract_name(text):

    lines = text.splitlines()

    for line in lines:

        s = " ".join(line.split()).strip()

        if not s:
            continue

        if len(s) > 60:
            continue

        if any(x in s for x in ["@", "http", "|", ":"]):
            continue

        if re.fullmatch(
            r"[A-Za-z][A-Za-z .'-]{2,59}",
            s
        ):

            words = s.split()

            if 2 <= len(words) <= 5:
                return s

    return ""


# --------------------------------------------------
# LINKS
# --------------------------------------------------

def extract_links(text):

    urls = re.findall(
        r"https?://[^\s)\]>]+",
        text,
        re.I
    )

    linkedin = next(
        (url for url in urls if "linkedin.com" in url.lower()),
        ""
    )

    github = next(
        (url for url in urls if "github.com" in url.lower()),
        ""
    )

    return {
        "linkedin": linkedin,
        "github": github
    }


# --------------------------------------------------
# SKILLS
# --------------------------------------------------

def extract_skills(text):

    low = text.lower()

    found = []

    for skill in SKILL_KEYWORDS:

        if skill.lower() in low:
            found.append(skill)

    return found


# --------------------------------------------------
# SECTION EXTRACTION
# --------------------------------------------------

def extract_section(text, headings):

    lines = text.splitlines()

    start = None

    for i, line in enumerate(lines):

        clean = line.strip()

        for heading in headings:

            if re.fullmatch(
                rf"\s*{re.escape(heading)}\s*:?",
                clean,
                re.I
            ):

                start = i + 1
                break

        if start is not None:
            break

    if start is None:
        return ""

    stops = [
        "skills",
        "technical skills",
        "education",
        "experience",
        "work experience",
        "professional experience",
        "projects",
        "certifications",
        "summary",
        "profile",
        "achievements",
        "contact"
    ]

    output = []

    for line in lines[start:]:

        clean = line.strip()

        if any(
            re.fullmatch(
                rf"\s*{re.escape(stop)}\s*:?",
                clean,
                re.I
            )
            for stop in stops
        ):
            break

        if clean:
            output.append(clean)

    return "\n".join(output[:30])


# --------------------------------------------------
# MAIN PARSER
# --------------------------------------------------

def parse_resume(path):

    text = extract_text(path)

    links = extract_links(text)

    return {

        "name": extract_name(text),

        "email": extract_email(text),

        "phone": extract_phone(text),

        "linkedin": links["linkedin"],

        "github": links["github"],

        "skills": extract_skills(text),

        "education": extract_section(
            text,
            ["Education", "Academic Background"]
        ),

        "experience": extract_section(
            text,
            [
                "Experience",
                "Work Experience",
                "Professional Experience"
            ]
        ),

        "text_preview": text[:2000]
    }